"""python-docx helpers implementing the LPU report conventions (A4, TNR 12, 1.5 spacing, real Word fields)."""
import re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"
TEXT_WIDTH_CM = 15.5


def _set_font(style_or_run, name=FONT):
    style_or_run.font.name = name
    rpr = style_or_run.element.get_or_add_rPr() if hasattr(style_or_run, "element") and hasattr(style_or_run.element, "get_or_add_rPr") else style_or_run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(a), name)


class Report:
    def __init__(self, page_map=None):
        self.doc = Document()
        self.page_map = page_map or {}
        self.headings = []      # (level, text)
        self.figures = []       # caption strings "Figure N: ..."
        self.tables = []
        self.fig_n = 0
        self.tab_n = 0
        self.eq_n = 0
        self._styles()
        self._page_setup()

    # ───────────── styles / page ─────────────
    def _styles(self):
        st = self.doc.styles
        n = st["Normal"]; _set_font(n); n.font.size = Pt(12)
        pf = n.paragraph_format; pf.line_spacing = 1.5; pf.space_after = Pt(6); pf.space_before = Pt(0)
        n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for name, size, before, after, italic, pb in (("Heading 1", 15, 0, 14, False, True), ("Heading 2", 13, 14, 6, False, False), ("Heading 3", 12, 10, 4, True, False)):
            h = st[name]; _set_font(h); h.font.size = Pt(size); h.font.bold = True; h.font.italic = italic
            h.font.color.rgb = RGBColor(0, 0, 0)
            h.paragraph_format.space_before = Pt(before); h.paragraph_format.space_after = Pt(after)
            h.paragraph_format.keep_with_next = True; h.paragraph_format.line_spacing = 1.5
            h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT; h.paragraph_format.page_break_before = pb
        cap = st["Caption"]; _set_font(cap); cap.font.size = Pt(12); cap.font.bold = True; cap.font.italic = False
        cap.font.color.rgb = RGBColor(0, 0, 0); cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_before = Pt(4); cap.paragraph_format.space_after = Pt(10); cap.paragraph_format.line_spacing = 1.15
        for nm in ("List Bullet", "List Number"):
            s = st[nm]; _set_font(s); s.font.size = Pt(12); s.paragraph_format.line_spacing = 1.5; s.paragraph_format.space_after = Pt(3)
        for nm, left in (("TOC 1", 0), ("TOC 2", 0.6), ("TOC 3", 1.2), ("Table of Figures", 0)):
            s = st.add_style(nm, WD_STYLE_TYPE.PARAGRAPH); s.base_style = st["Normal"]; _set_font(s); s.font.size = Pt(12)
            s.paragraph_format.left_indent = Cm(left); s.paragraph_format.line_spacing = 1.15; s.paragraph_format.space_after = Pt(3)
            s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            s.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        s = st.add_style("CodeBlock", WD_STYLE_TYPE.PARAGRAPH); s.base_style = st["Normal"]; _set_font(s, "Courier New"); s.font.size = Pt(9)
        s.paragraph_format.line_spacing = 1.0; s.paragraph_format.space_after = Pt(0); s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        s.paragraph_format.left_indent = Cm(0.3)
        s = st.add_style("FrontHeading", WD_STYLE_TYPE.PARAGRAPH); s.base_style = st["Normal"]; _set_font(s); s.font.size = Pt(14); s.font.bold = True
        s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; s.paragraph_format.space_after = Pt(14)

    def _page_setup(self):
        s = self.doc.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(3.0), Cm(2.5), Cm(2.5), Cm(2.5)
        s.footer_distance = Cm(1.2)

    def _set_pgnum(self, section, fmt, start=1):
        sp = section._sectPr
        for e in sp.findall(qn("w:pgNumType")): sp.remove(e)
        pg = OxmlElement("w:pgNumType"); pg.set(qn("w:fmt"), fmt); pg.set(qn("w:start"), str(start))
        cols = sp.find(qn("w:cols"))
        if cols is not None: cols.addprevious(pg)          # schema order: pgSz, pgMar, ..., pgNumType, cols, ...
        else: sp.find(qn("w:pgMar")).addnext(pg)

    def _footer_pagenum(self, section):
        section.footer.is_linked_to_previous = False
        p = section.footer.paragraphs[0]
        for r in list(p.runs): r._r.getparent().remove(r._r)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.field(p, "PAGE", "1", size=11)

    def _footer_empty(self, section):
        section.footer.is_linked_to_previous = False
        for r in list(section.footer.paragraphs[0].runs): r._r.getparent().remove(r._r)

    def new_section(self, fmt=None, start=1, footer=True):
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.5)
        if fmt: self._set_pgnum(sec, fmt, start)
        (self._footer_pagenum if footer else self._footer_empty)(sec)
        return sec

    # ───────────── low-level text ─────────────
    def field(self, par, instr, cached="", size=None, bold=None):
        def run_with(child):
            r = par.add_run(); 
            if size: r.font.size = Pt(size)
            if bold is not None: r.font.bold = bold
            r._r.append(child); return r
        b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); run_with(b)
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; run_with(it)
        s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate"); run_with(s)
        r = par.add_run(cached); 
        if size: r.font.size = Pt(size)
        if bold is not None: r.font.bold = bold
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); run_with(e)

    TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|_\{[^}]*\}|\^\{[^}]*\}|\[\[.+?\]\])")

    def inline(self, par, text, size=None, bold=False, italic=False):
        """Mini-markup: **bold**, *italic*, _{sub}, ^{sup}, [[placeholder]] (yellow highlight)."""
        for part in self.TOKEN.split(text):
            if not part: continue
            b, i = bool(bold), bool(italic)
            if part.startswith("**") and part.endswith("**"): r = par.add_run(part[2:-2]); b = True
            elif part.startswith("[[") and part.endswith("]]"): r = par.add_run("[" + part[2:-2] + "]"); r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            elif part.startswith("*") and part.endswith("*") and len(part) > 2: r = par.add_run(part[1:-1]); i = True
            elif part.startswith("_{"): r = par.add_run(part[2:-1]); r.font.subscript = True
            elif part.startswith("^{"): r = par.add_run(part[2:-1]); r.font.superscript = True
            else: r = par.add_run(part)
            if b: r.bold = True          # only set when True so that style-level bold/italic (headings) is not overridden
            if i: r.italic = True
            if size: r.font.size = Pt(size)
        return par

    # ───────────── block elements ─────────────
    def p(self, text, align=None, size=None, bold=False, italic=False, space_after=None, keep_next=False, indent_first=False, style=None):
        par = self.doc.add_paragraph(style=style)
        self.inline(par, text, size=size, bold=bold, italic=italic)
        if align == "center": par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "left": par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif align == "right": par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if space_after is not None: par.paragraph_format.space_after = Pt(space_after)
        if keep_next: par.paragraph_format.keep_with_next = True
        if indent_first: par.paragraph_format.first_line_indent = Cm(1.0)
        return par

    def blank(self, n=1, size=12):
        for _ in range(n):
            par = self.doc.add_paragraph(); par.paragraph_format.space_after = Pt(0)
            par.add_run("").font.size = Pt(size)

    def bullets(self, items, numbered=False):
        for it in items:
            par = self.doc.add_paragraph(style="List Number" if numbered else "List Bullet"); self.inline(par, it)
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    def h(self, level, text):
        par = self.doc.add_heading(level=level); self.inline(par, text)
        for r in par.runs: r.font.color.rgb = RGBColor(0, 0, 0); _set_font(r)
        self.headings.append((level, text)); return par

    def equation(self, text, number=True):
        par = self.doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = par.paragraph_format; pf.line_spacing = 1.3; pf.space_after = Pt(6); pf.space_before = Pt(2)
        pf.tab_stops.add_tab_stop(Cm(TEXT_WIDTH_CM / 2), WD_TAB_ALIGNMENT.CENTER); pf.tab_stops.add_tab_stop(Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
        par.add_run("\t"); self.inline(par, text, italic=True)
        if number:
            self.eq_n += 1; par.add_run(f"\t({self.eq_n})")
        return par

    def code(self, text):
        for line in text.strip("\n").split("\n"):
            par = self.doc.add_paragraph(style="CodeBlock"); par.add_run(line)
            ppr = par._p.get_or_add_pPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "F2F2F2"); ppr.append(shd)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def caption(self, kind, text):
        par = self.doc.add_paragraph(style="Caption")
        if kind == "Figure": self.fig_n += 1; n = self.fig_n
        else: self.tab_n += 1; n = self.tab_n
        par.add_run(f"{kind} ")
        self.field(par, f"SEQ {kind} \\* ARABIC", str(n), bold=True)
        par.add_run(f": {text}")
        (self.figures if kind == "Figure" else self.tables).append(f"{kind} {n}: {text}")
        if kind == "Table": par.paragraph_format.keep_with_next = True
        return par

    def figure(self, path, caption, width_cm=15.0):
        par = self.doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.keep_with_next = True; par.paragraph_format.space_after = Pt(2); par.paragraph_format.line_spacing = 1.0
        par.add_run().add_picture(str(path), width=Cm(width_cm))
        self.caption("Figure", caption)

    def table(self, header, rows, caption, widths=None, size=10.5, align=None, bold_first_col=False, highlight_row=None):
        self.caption("Table", caption)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        widths = widths or [TEXT_WIDTH_CM / len(header)] * len(header)
        align = align or ["left"] + ["center"] * (len(header) - 1)
        def fill(cell, txt, bold=False, shade=None, al="left"):
            cell.text = ""; par = cell.paragraphs[0]; par.paragraph_format.line_spacing = 1.0; par.paragraph_format.space_after = Pt(1); par.paragraph_format.space_before = Pt(1)
            par.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}[al]
            self.inline(par, str(txt), size=size, bold=bold)
            if shade:
                tcpr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), shade); tcpr.append(shd)
        for j, hd in enumerate(header): fill(t.rows[0].cells[j], hd, bold=True, shade="D9D9D9", al="center")
        for i, row in enumerate(rows):
            for j, v in enumerate(row):
                fill(t.rows[i + 1].cells[j], v, bold=(bold_first_col and j == 0) or (highlight_row == i), shade="E2EFDA" if highlight_row == i else None, al=align[j])
        t.autofit = False        # python-docx inserts <w:tblLayout w:type="fixed"/> in the schema-correct position
        for j, w_ in enumerate(widths): t.columns[j].width = Cm(w_)
        for row in t.rows:
            for j, c in enumerate(row.cells): c.width = Cm(widths[j])
        # repeat header row & avoid row splitting
        trpr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trpr.append(th)
        for row in t.rows:
            trp = row._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trp.append(cs)
        for row in t.rows[:-1]:                      # keep the whole table on one page when it fits
            for c in row.cells:
                for par in c.paragraphs: par.paragraph_format.keep_with_next = True
        sp = self.doc.add_paragraph(); sp.paragraph_format.space_after = Pt(6); sp.paragraph_format.line_spacing = 1.0
        return t

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ───────────── Word-field tables of contents ─────────────
    def toc_block(self, anchor, instr, entries, style_for):
        """Insert a multi-paragraph TOC-style field before `anchor`. entries = [(level, text, page_str)]."""
        paras = []
        for i, (lvl, text, page) in enumerate(entries):
            par = self.doc.add_paragraph(style=style_for(lvl))
            if i == 0:
                for typ, txt in (("begin", None), ("instr", instr), ("separate", None)):
                    r = par.add_run()
                    if typ == "instr":
                        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {txt} "; r._r.append(it)
                    else:
                        fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), typ); r._r.append(fc)
            par.add_run(text); par.add_run("\t" + str(page))
            if i == len(entries) - 1:
                r = par.add_run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "end"); r._r.append(fc)
            paras.append(par)
        for par in paras: anchor._p.addprevious(par._p)
        anchor._p.getparent().remove(anchor._p)

    def save(self, path):
        z = self.doc.settings.element.find(qn("w:zoom"))
        if z is not None and z.get(qn("w:percent")) is None: z.set(qn("w:percent"), "100")
        self.doc.save(str(path))
