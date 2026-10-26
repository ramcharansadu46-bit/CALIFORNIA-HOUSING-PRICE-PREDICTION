"""Cover page, declaration, certificate, acknowledgement (LPU template conventions)."""
from pathlib import Path
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx_helpers import _set_font
from rc import TITLE, TITLE_MAIN, TITLE_SUB, STUDENT


def front_heading(rp, text):
    """Declaration / Certificate / Acknowledgement headings: 12 pt, bold, underlined, centred (as in the LPU template)."""
    par = rp.p(text, align="center", bold=True, size=12, space_after=14)
    for r in par.runs: r.underline = True
    return par


NAVY = RGBColor(0x1F, 0x38, 0x64)          # title colour measured from the reference LPU report cover
LOGO = Path(__file__).parent / "assets" / "lpu_logo.jpg"
# space-before (pt) of each cover line; tuned numerically against the reference cover (see tools notes)
COVER_BEFORE = {"logo": 0.7, "project": 42.2, "duration": 7.7, "title": 50.1, "sub1": 10.6, "sub2": 0, "submitted": 81.5, "name": 3.8,
                "reg": 0.0, "prog": 0.0, "course": 0.5, "under": 26.7, "guide": 3.9, "desig": 0.0, "discipline": 1.6, "school": 0.0, "univ": 0.0, "punjab": 0.0}


def cover(rp):
    """Cover page in the style of the LPU project report: logo, navy title, Calibri, centred."""
    d = rp.doc
    sec = d.sections[0]; sec.left_margin = sec.right_margin = Cm(2.5)        # symmetric margins -> truly centred
    B = COVER_BEFORE

    def line(key, text, size=11, bold=False, italic=False, color=None, pitch=None):
        par = d.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = par.paragraph_format; pf.space_before = Pt(B[key]); pf.space_after = Pt(0)
        if pitch:
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY; pf.line_spacing = Pt(pitch)
        else:
            pf.line_spacing = 1.0
        r = par.add_run(text); r.font.size = Pt(size); r.bold = bold or None; r.italic = italic or None
        _set_font(r, "Calibri")
        if color is not None: r.font.color.rgb = color
        return par

    logo = d.add_paragraph(); logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo.paragraph_format.space_before = Pt(B["logo"]); logo.paragraph_format.space_after = Pt(0); logo.paragraph_format.line_spacing = 1.0
    logo.add_run().add_picture(str(LOGO), width=Cm(9.41))
    line("project", "PROJECT REPORT", 16, bold=True, pitch=20)
    line("duration", "(Project Semester: August – December 2026)", 11, italic=True, pitch=18)
    line("title", TITLE_MAIN, 20, bold=True, color=NAVY, pitch=26)
    sub = line("sub1", TITLE_SUB, 12, italic=True, pitch=18)
    line("submitted", "Submitted by", pitch=18.7)
    line("name", STUDENT, 13, bold=True, pitch=19)
    line("reg", "Registration No.: 12502844", pitch=18.7)
    line("prog", "Programme and Section: Bachelor of Technology in Computer Science and Engineering & K2P24QR", pitch=18.7)
    line("course", "Course Code: INT234", pitch=18.7)
    line("under", "Under the Guidance of", pitch=18.7)
    line("guide", "Dr. Mrinalini Rana", 13, bold=True, pitch=19)
    line("desig", "Assistant Professor (U.Id 22138)", pitch=18.7)
    line("discipline", "Discipline of CSE/IT", pitch=18.7)
    line("school", "Lovely School of Computer Science and Engineering", bold=True, pitch=18.7)
    line("univ", "Lovely Professional University, Phagwara", bold=True, pitch=18.7)
    line("punjab", "Punjab, India – 144411", pitch=18.7)


def declaration(rp):
    front_heading(rp, "DECLARATION")
    rp.p(f"I, {STUDENT}, student of Bachelor of Technology in Computer Science and Engineering under CSE/IT Discipline at Lovely Professional University, Punjab, hereby declare that all the information furnished in this project report is based on my own intensive work and is genuine.")
    rp.blank(3)
    rp.p("Date: 31-10-2026", align="left", space_after=0)
    rp.p(f"Signature: {STUDENT}", align="right", space_after=0)
    rp.p("Registration No. 12502844", align="left", space_after=0)
    rp.p(f"Name of the student: {STUDENT}", align="right")


def certificate(rp):
    front_heading(rp, "CERTIFICATE")
    rp.p(f"This is to certify that {STUDENT} bearing Registration no. 12502844 has completed INT234 project titled, **“{TITLE}”** under my guidance and supervision. To the best of my knowledge, the present work is the result of his/her original development, effort and study.")
    rp.blank(3)
    rp.p("Dr. Mrinalini Rana", align="left", bold=True, space_after=0)
    rp.p("Assistant Professor (U.Id 22138)", align="left", bold=True, space_after=0)
    rp.p("Lovely School of Computer Science and Engineering", align="left", bold=True, space_after=0)
    rp.p("Lovely Professional University", align="left", space_after=0)
    rp.p("Phagwara, Punjab.", align="left", space_after=0)
    rp.blank(1)
    rp.p("Date: 31-10-2026", align="left")


def acknowledgement(rp):
    front_heading(rp, "ACKNOWLEDGEMENT")
    rp.p("I express my sincere gratitude to my faculty coordinator, Dr. Mrinalini Rana, Lovely Professional University, for the guidance, encouragement and constructive feedback that shaped this project from the choice of problem to the final analysis. I also thank the School and the Discipline of CSE/IT for providing the academic environment and the course framework in which this work was carried out.")
    rp.p("I acknowledge the original authors of the dataset, R. Kelley Pace and Ronald Barry, whose 1997 work made the California census data available through the StatLib repository, and the contributors who curated and re-published it on Kaggle and OpenML, as well as Aurélien Géron, whose public repository provided the copy used here. The project would not have been possible without the open-source scientific Python ecosystem: NumPy, pandas, SciPy, statsmodels, scikit-learn, XGBoost, SHAP, Matplotlib and Seaborn.")
    rp.p("Finally, I thank my family and friends for their patience and support during the preparation of this report.")
    rp.blank(2)
    rp.p(STUDENT, align="right", bold=True, space_after=0)
    rp.p("Date: 31-10-2026", align="right")
